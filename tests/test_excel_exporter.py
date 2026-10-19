"""
Tests — ExcelExporter.

Teste l'export Excel : création de fichier, formatage, contenu,
et performance (< 5s pour 100 enregistrements).
"""

import os
import sys
import time
from datetime import date

import pytest

ROOT_DIR = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, ROOT_DIR)

from export.excel_exporter import ExcelExporter


class TestExcelExporter:
    """Tests pour ExcelExporter."""

    @pytest.fixture
    def exporter(self, db_manager, tmp_exports_dir):
        """Fournit une instance ExcelExporter avec dossier temporaire."""
        exp = ExcelExporter(db_manager)
        exp._exports_dir = tmp_exports_dir
        return exp

    def test_export_creates_file(self, exporter):
        """L'export crée un fichier .xlsx."""
        filepath = exporter.export_today()
        assert os.path.exists(filepath)
        assert filepath.endswith(".xlsx")

    def test_export_filename_format(self, exporter):
        """Le nom du fichier contient la date du jour."""
        filepath = exporter.export_today()
        today = date.today().strftime("%Y-%m-%d")
        assert today in os.path.basename(filepath)

    def test_export_with_records(self, exporter, db_manager, sample_employee):
        """L'export contient les enregistrements de présence."""
        db_manager.add_attendance_record(
            employee_id=sample_employee["id"],
            status="IN_TIME",
            confidence=0.95,
        )
        filepath = exporter.export_today()
        assert os.path.exists(filepath)
        # Vérifier que le fichier n'est pas vide
        assert os.path.getsize(filepath) > 0

    def test_export_with_mixed_statuses(self, exporter, db_manager):
        """L'export gère les statuts IN_TIME et LATE."""
        emp1_id = db_manager.add_employee("Early Bird", "HR")
        emp2_id = db_manager.add_employee("Late Comer", "IT")
        db_manager.add_attendance_record(emp1_id, "IN_TIME", 0.92)
        db_manager.add_attendance_record(emp2_id, "LATE", 0.88)

        filepath = exporter.export_today()
        assert os.path.exists(filepath)

    def test_export_empty_records(self, exporter):
        """L'export fonctionne sans enregistrements."""
        filepath = exporter.export_today()
        assert os.path.exists(filepath)

    def test_export_performance_100_records(self, exporter, db_manager):
        """L'export de 100 enregistrements prend < 5 secondes."""
        # Créer 100 enregistrements
        for i in range(100):
            emp_id = db_manager.add_employee(f"Perf User {i}", "Test")
            status = "IN_TIME" if i % 2 == 0 else "LATE"
            db_manager.add_attendance_record(emp_id, status, 0.9)

        start = time.perf_counter()
        filepath = exporter.export_today()
        elapsed = time.perf_counter() - start

        assert os.path.exists(filepath)
        assert elapsed < 5.0, f"Export trop lent : {elapsed:.2f}s (max 5s)"

    def test_export_xlsx_content(self, exporter, db_manager, sample_employee):
        """Le fichier Excel contient les bonnes colonnes."""
        from openpyxl import load_workbook

        db_manager.add_attendance_record(
            employee_id=sample_employee["id"],
            status="IN_TIME",
            confidence=0.95,
        )
        filepath = exporter.export_today()

        wb = load_workbook(filepath)
        ws = wb.active
        # Vérifier le titre de la feuille
        assert "Présences" in ws.title
        # Vérifier les en-têtes (ligne 3)
        headers = [ws.cell(row=3, column=c).value for c in range(1, 6)]
        assert "Nom" in headers
        assert "Statut" in headers
        assert "Classe" in headers
